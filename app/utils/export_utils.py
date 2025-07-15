"""
Utility functions for exporting data to various formats
"""
import os
import csv
import json
import tempfile
import logging
from typing import List, Dict, Any, Optional, Tuple

# Setup logging
logger = logging.getLogger(__name__)
logger.setLevel(logging.DEBUG)

def generate_csv_file(
    data: List[Dict[str, Any]], 
    custom_headers: Optional[str] = None
) -> Tuple[str, str]:
    """
    Generate a CSV file from data
    
    Args:
        data: List of dictionaries containing data to export
        custom_headers: Optional comma-separated custom headers
        
    Returns:
        Tuple containing:
            - Path to the generated file
            - Filename for the download
    """
    try:
        # Create temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.csv') as temp_file:
            file_path = temp_file.name
        
        # Get headers
        if custom_headers:
            # Use custom headers if provided
            headers = [h.strip() for h in custom_headers.split(',')]
        else:
            # Get all possible headers from data
            all_headers = set()
            for record in data:
                all_headers.update(record.keys())
            
            # Sort headers to ensure 'source_pdf' and 'extraction_date' come first
            headers = ['source_pdf', 'extraction_date']
            for header in sorted(all_headers):
                if header not in headers:
                    headers.append(header)
        
        # Write CSV
        with open(file_path, 'w') as f:
            writer = csv.DictWriter(f, fieldnames=headers, extrasaction='ignore')
            writer.writeheader()
            
            for record in data:
                writer.writerow({k: str(v) if v is not None else "" for k, v in record.items()})
        
        logger.debug(f"CSV file created at {file_path}")
        return file_path, 'bank_details.csv'
    
    except Exception as e:
        logger.error(f"Error generating CSV file: {str(e)}")
        raise

def generate_excel_file(
    data: List[Dict[str, Any]], 
    custom_headers: Optional[str] = None
) -> Tuple[str, str]:
    """
    Generate an Excel file from data
    
    Args:
        data: List of dictionaries containing data to export
        custom_headers: Optional comma-separated custom headers
        
    Returns:
        Tuple containing:
            - Path to the generated file
            - Filename for the download
    """
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        
        # Create temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as temp_file:
            file_path = temp_file.name
        
        # Get headers
        if custom_headers:
            # Use custom headers if provided
            headers = [h.strip() for h in custom_headers.split(',')]
        else:
            # Get all possible headers from data
            all_headers = set()
            for record in data:
                all_headers.update(record.keys())
            
            # Sort headers to ensure 'source_pdf' and 'extraction_date' come first
            headers = ['source_pdf', 'extraction_date']
            for header in sorted(all_headers):
                if header not in headers:
                    headers.append(header)
        
        # Create Excel file
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Bank Details"
        
        # Add headers
        for col_idx, header in enumerate(headers, 1):
            ws.cell(row=1, column=col_idx, value=header)
        
        # Add data
        for row_idx, record in enumerate(data, 2):
            for col_idx, header in enumerate(headers, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(record.get(header, "")))
        
        # Auto-adjust column width
        for col_idx, header in enumerate(headers, 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = max(len(header) + 2, 15)
        
        # Save Excel file
        wb.save(file_path)
        
        logger.debug(f"Excel file created at {file_path}")
        return file_path, 'bank_details.xlsx'
    
    except Exception as e:
        logger.error(f"Error generating Excel file: {str(e)}")
        raise

def generate_json_file(data: List[Dict[str, Any]]) -> Tuple[str, str]:
    """
    Generate a JSON file from data
    
    Args:
        data: List of dictionaries containing data to export
        
    Returns:
        Tuple containing:
            - Path to the generated file
            - Filename for the download
    """
    try:
        # Create temporary file
        with tempfile.NamedTemporaryFile(delete=False, suffix='.json') as temp_file:
            file_path = temp_file.name
        
        # Write JSON file
        with open(file_path, 'w') as f:
            json.dump({"bank_details": data}, f, indent=2)
        
        logger.debug(f"JSON file created at {file_path}")
        return file_path, 'bank_details.json'
    
    except Exception as e:
        logger.error(f"Error generating JSON file: {str(e)}")
        raise 